"""Build docs/belief_methods/Cube_push_WAFR08.pptx: the MuJoCo cube push with WAFR08 and its belief renders
(uploaded to Google Slides afterwards; the mp4s are embedded there from Drive).

    MUJOCO_GL=egl .venv/bin/python examples/contact_rich_uncertainty/cube_push_wafr08.py --seed {0,1,2}
    .venv/bin/python docs/belief_methods/build_cube_push_wafr08_slides.py
"""
import json
import os
import subprocess

import imageio_ffmpeg
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

plt.rcParams.update({"text.usetex": True, "text.latex.preamble": r"\usepackage{amsmath,amssymb,bm}"})

HERE = os.path.dirname(os.path.abspath(__file__))
RUN = os.path.join(HERE, "..", "..", "runs", "cube_push_wafr08")
IMG = os.path.join(HERE, "cube_push")
os.makedirs(IMG, exist_ok=True)
SEEDS = (0, 1, 2)
FRACTIONS = (0.02, 0.3, 0.6, 0.995)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
GREY = RGBColor(0x55, 0x55, 0x55)


def title_only(title, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    t = s.shapes.title
    t.text = title
    t.text_frame.paragraphs[0].font.size = Pt(28)
    t.left, t.top, t.width, t.height = Inches(0.4), Inches(0.15), W - Inches(0.8), Inches(0.8)
    if subtitle:
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.9), W - Inches(1), Inches(0.4))
        tb.text_frame.text = subtitle
        tb.text_frame.paragraphs[0].font.size = Pt(14)
        tb.text_frame.paragraphs[0].font.color.rgb = GREY
    return s


def text(s, x, y, w, h, lines, size=15):
    tf = s.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        level = (len(line) - len(line.lstrip(" "))) // 2
        p.text, p.level = line.strip(), level
        p.font.size = Pt(size - 2 * level)
    return tf


def latex(name, tex, size=20):
    path = os.path.join(IMG, f"{name}.png")
    fig = plt.figure()
    fig.text(0, 0, tex, fontsize=size)
    fig.savefig(path, bbox_inches="tight", pad_inches=0.05, dpi=200, transparent=True)
    plt.close(fig)
    return path


def keyframes(seed):
    """Frames at fixed fractions of the video, as jpg."""
    video = os.path.join(RUN, "videos", f"wafr08_seed{seed}.mp4")
    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    n = imageio_ffmpeg.count_frames_and_secs(video)[1]
    paths = []
    for k, f in enumerate(FRACTIONS):
        path = os.path.join(IMG, f"seed{seed}_{k}.jpg")
        subprocess.run([ffmpeg, "-y", "-loglevel", "error", "-ss", f"{f * n:.2f}", "-i", video, "-frames:v", "1",
                        "-q:v", "3", path], check=True)
        paths.append(path)
    return paths


def evolution_plot(results):
    path = os.path.join(IMG, "evolution.png")
    plt.rcParams["text.usetex"] = False
    fig, axes = plt.subplots(1, 3, figsize=(12.5, 3.6), dpi=130)
    for seed, r in results.items():
        log = r["log"]
        k = range(1, len(log) + 1)
        goal = [i + 1 for i, l in enumerate(log) if l["wrt"].startswith("goal")]
        axes[0].plot(k, [l["entropy"] for l in log], ".-", label=f"seed {seed}")
        axes[1].plot(k, [l["p_goal"] for l in log], ".-", label=f"seed {seed}")
        axes[2].plot(k, [100 * l["goal_dist"] for l in log], ".-", label=f"seed {seed}")
        axes[2].plot(goal, [100 * log[i - 1]["goal_dist"] for i in goal], "k^", ms=5)
    axes[0].set_title("belief entropy H(b) [nats]")
    axes[1].set_title("P_b(cube within goal tolerance)")
    axes[1].axhline(0.8, color="g", ls="--", lw=1, label="stop: 1 - δ")
    axes[2].set_title("true distance to goal [cm] (▲ goal push)")
    axes[2].axhline(2.0, color="g", ls="--", lw=1, label="tolerance")
    for a in axes:
        a.set_xlabel("WRT #")
        a.legend(fontsize=8)
        a.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    plt.rcParams["text.usetex"] = True
    return path


results = {s: json.load(open(os.path.join(RUN, f"seed{s}.json"))) for s in SEEDS}

# ---------------------------------------------------------------- title
s = title_only("Cube push with WAFR08 belief-based execution (MuJoCo, skrl)",
               "Hidden cube pose + friction μ, mass m, damping k_d; the only command is the wanted final cube pose")
text(s, Inches(0.6), Inches(1.6), Inches(6.2), Inches(5.5), [
    "Setup",
    "  Pusher: cylinder, position targets (x, y, z), sees its own position and actuator force (noisy)",
    "  Cube: hidden start pose (x, y, yaw mod 90°) and μ ∈ [0.3, 0.8], m ∈ [0.2, 0.6] kg, k_d ∈ [0, 3] N s/m",
    "  Goal: cube pose within 2 cm and 10°",
    "Planner: WAFR08 (Hsiao, Lozano-Pérez, Kaelbling 2008)",
    "  Grid belief over (x, y, yaw, μ, m, k_d): 36 × 35 × 9 × 3 × 3 × 3 cells",
    "  World-relative trajectories at the MAP state w*: guarded touch, force probe, goal push",
    "  Every WRT simulated from every belief hypothesis with batched mujoco.rollout",
    "Result: 3 / 3 seeds reach the goal (21–27 WRTs, 35–46 s simulated)",
], size=16)
s.shapes.add_picture(os.path.join(RUN, "videos", "wafr08_seed1_final.png"), Inches(7.0), Inches(1.5),
                     width=Inches(5.9))

# ---------------------------------------------------------------- method
s = title_only("WAFR08 decision rule on the cube push", "Belief update with each observed contact, travel and force")
s.shapes.add_picture(latex("rule", "\n".join([
    r"$b'(w') \propto \sum_{w:\,T_\tau(w)=w'} P(o \mid w, \tau)\, b(w)$ \quad (grid, multilinear splat of $T_\tau(w)$)",
    r"$o = (c,\ \phi,\ \bar F_1,\ \bar F_2,\ F_{\rm peak})$, \quad $\bar F \approx \mu m g + k_d v$, \quad "
    r"$F_{\rm peak} \approx m a$",
    r"1. stop if $P_b(\mathrm{goal}) \geq 1-\delta$ \quad ($\delta = 0.2$)",
    r"2. run goal push $\tau_g(w^*)$ if $P_b\big(w \in G_\delta(w^*)\big) \geq 1-\delta$ \quad (belief pre-image)",
    r"3. else $\tau^* = \arg\max_{\tau}\ IG(\tau, b) - c\, FP(\tau, b)$, \quad "
    r"$IG = H(b) - \sum_j b(w_j)\, H\big(b'_\tau \mid o_j\big)$",
])), Inches(0.6), Inches(1.5), width=Inches(12.0))
text(s, Inches(0.6), Inches(5.4), Inches(12.2), Inches(1.9), [
    "Changes vs. the paper: IG is measured on the next-state belief b'_τ (actions that scatter the cube lose what they "
    "disturb); G_δ uses 2× tolerance for waypoint pushes that do not end at the goal; force probes only once the pose "
    "entropy < 1.5 nats; FP = P(pusher lands on the cube) + P(cube leaves the table).",
], size=14)

# ---------------------------------------------------------------- per-seed keyframes
for seed in SEEDS:
    r = results[seed]
    kinds = [l["wrt"].split()[0] for l in r["log"]]
    mu, m, kd = r["physics"]
    s = title_only(f"Seed {seed}: belief over time (keyframes)",
                   f"true μ {mu:.2f}, m {m:.2f} kg, k_d {kd:.2f} | {len(kinds)} WRTs: {kinds.count('touch')} touches, "
                   f"{kinds.count('probe')} probes, {kinds.count('goal')} goal pushes | final error "
                   f"{r['goal_dist_cm']:.1f} cm, {r['goal_dyaw_deg']:.1f}°")
    for k, path in enumerate(keyframes(seed)):
        x = Inches(0.35 + 6.35 * (k % 2))
        y = Inches(1.35 + 3.05 * (k // 2))
        s.shapes.add_picture(path, x, y, height=Inches(3.0))
    text(s, Inches(4.15), Inches(1.35), Inches(2.4), Inches(5.9), [
        "Panels: overview, top-down, pusher follow, belief plots",
        "Orange tiles: most likely cube cells (opacity ∝ mass)",
        "Blue box: MAP w*",
        "Green box + disc: goal pose, 2 cm tolerance",
        "Red: table limit",
        "Yellow: current WRT path",
        "Order: prior → touches → goal pushes → stop",
    ], size=11)

# ---------------------------------------------------------------- video slides (videos embedded from Drive)
for seed in SEEDS:
    title_only(f"Seed {seed}: full render", f"Drive video wafr08_seed{seed}.mp4 "
               f"(runs/cube_push_wafr08/videos/)")

# ---------------------------------------------------------------- results
s = title_only("Belief evolution and results", "3 / 3 seeds reach the 2 cm / 10° goal; WAFR08 stops on belief, "
               "and the true cube is inside the tolerance too")
s.shapes.add_picture(evolution_plot(results), Inches(0.3), Inches(1.4), width=Inches(12.7))
rows = [f"seed {s_}: {len(r['log'])} WRTs, {r['sim_time_s']:.0f} s simulated, error {r['goal_dist_cm']:.1f} cm "
        f"{r['goal_dyaw_deg']:.1f}°, true (μ, m, k_d) = ({r['physics'][0]:.2f}, {r['physics'][1]:.2f}, "
        f"{r['physics'][2]:.2f}), MAP levels ({r['log'][-1]['map'][3]:.2f}, {r['log'][-1]['map'][4]:.2f}, "
        f"{r['log'][-1]['map'][5]:.1f})" for s_, r in results.items()]
text(s, Inches(0.6), Inches(5.2), Inches(12.2), Inches(2.0), rows, size=13)

# ---------------------------------------------------------------- limits
s = title_only("What works, what does not yet")
text(s, Inches(0.6), Inches(1.3), Inches(12.2), Inches(5.9), [
    "Works",
    "  Guarded touches localize the cube from a 12 × 20 cm, any-yaw prior in ~10 WRTs",
    "  Goal pushes are executed only from the belief pre-image; touches re-localize between them",
    "  Hypothesis rollouts reproduce the real sensor trace exactly (same MuJoCo model, same controls)",
    "Does not yet",
    "  Physics is only coarsely identified: 3 levels per parameter, and push forces mostly measure μ·m",
    "  Force probes almost never fire (gated on pose entropy); belief stays ~5 nats at the stop",
    "  ~2 min wall per episode on the laptop CPU (400–6000 hypotheses per WRT)",
    "Next",
    "  Audit, GenAIF, quasi-static belief dynamics and RSS20 on the same env, same video tooling",
    "Code: skrl/envs/cube_push/{sim,env,wafr08,render}.py, skrl/belief/grid_belief.py, "
    "examples/contact_rich_uncertainty/cube_push_wafr08.py (commit dc5467f5)",
], size=17)

out = os.path.join(HERE, "Cube_push_WAFR08.pptx")
prs.save(out)
print(out, len(prs.slides), "slides")
