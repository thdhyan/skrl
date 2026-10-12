"""Build docs/belief_methods/Belief_methods_equations.pptx (uploaded to Google Slides afterwards).

    .venv/bin/python docs/belief_methods/build_belief_methods_slides.py

Equations are rendered with LaTeX (matplotlib usetex) into docs/belief_methods/eq/*.png.
"""
import hashlib
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

plt.rcParams.update({"text.usetex": True, "text.latex.preamble": r"\usepackage{amsmath,amssymb,bm}"})

HERE = os.path.dirname(os.path.abspath(__file__))
EQ = os.path.join(HERE, "eq")
os.makedirs(EQ, exist_ok=True)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
GREY = RGBColor(0x55, 0x55, 0x55)


def render(latex, size=22):
    """LaTeX (display math, without $) -> cached transparent PNG path."""
    path = os.path.join(EQ, hashlib.md5(f"{size}{latex}".encode()).hexdigest()[:12] + ".png")
    if not os.path.exists(path):
        fig = plt.figure()
        fig.text(0, 0, rf"$\displaystyle {latex}$", fontsize=size)
        fig.savefig(path, bbox_inches="tight", pad_inches=0.03, dpi=220, transparent=True)
        plt.close(fig)
    return path


def title_only(title, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    t = s.shapes.title
    t.text = title
    t.text_frame.paragraphs[0].font.size = Pt(28)
    t.left, t.top, t.width, t.height = Inches(0.4), Inches(0.15), W - Inches(0.8), Inches(0.8)
    if subtitle:
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.9), W - Inches(1), Inches(0.4))
        tb.text_frame.text = subtitle
        tb.text_frame.paragraphs[0].font.size = Pt(14)
        tb.text_frame.paragraphs[0].font.color.rgb = GREY
    return s


def text(s, x, y, w, h, lines, size=15):
    tf = s.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
    return tf


def equations(title, rows, subtitle=None, note=None, size=22, cap_w=3.3):
    """rows: [(caption, latex)]. Captions on the left, equations stacked on the right, scaled to fit."""
    s = title_only(title, subtitle)
    top, bottom = Inches(1.4), H - (Inches(0.75) if note else Inches(0.25))
    x_eq, max_w = Inches(0.4 + cap_w + 0.2), W - Inches(0.4 + cap_w + 0.2 + 0.4)
    imgs = [(cap, render(tex, size)) for cap, tex in rows]
    px = [Image.open(p).size for _, p in imgs]
    gap = Inches(0.12)
    # native size at 220 dpi -> EMU, then one common scale so all rows fit
    native = [(Emu(int(w / 220 * 914400)), Emu(int(h / 220 * 914400))) for w, h in px]
    # caption height at 14 pt: ~11 characters per inch, 0.24 in per line
    cap_h = [Inches(0.24 * (1 + len(cap) // int(cap_w * 11))) if cap else 0 for cap, _ in imgs]
    scale = min(1.0, min(max_w / w for w, _ in native))
    while scale > 0.3 and sum(max(h * scale, c) for (_, h), c in zip(native, cap_h)) + gap * (len(imgs) - 1) > bottom - top:
        scale *= 0.95
    y = top
    for (cap, p), (w, h), c in zip(imgs, native, cap_h):
        w, h = int(w * scale), int(h * scale)
        if cap:
            tf = text(s, Inches(0.4), y, Inches(cap_w), max(h, c), [cap], size=14)
            tf.paragraphs[0].font.color.rgb = GREY
        s.shapes.add_picture(p, x_eq, y, w, h)
        y += max(h, c) + gap
    if note:
        text(s, Inches(0.5), H - Inches(0.7), W - Inches(1), Inches(0.6), [note], size=13)
    return s


def bullets(title, items, subtitle=None, size=18):
    s = title_only(title, subtitle)
    tf = text(s, Inches(0.6), Inches(1.4), W - Inches(1.2), H - Inches(1.7), [], size)
    for i, it in enumerate(items):
        level = (len(it) - len(it.lstrip(" "))) // 2
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.level = ("• " if level == 0 else "– ") + it.strip(), level
        p.font.size = Pt(size - 3 * level)
    return s


def table(title, header, rows, size=12, subtitle=None, widths=None):
    s = title_only(title, subtitle)
    shape = s.shapes.add_table(len(rows) + 1, len(header), Inches(0.3), Inches(1.4), W - Inches(0.6), Inches(0.4))
    if widths:
        for j, wj in enumerate(widths):
            shape.table.columns[j].width = Inches(wj)
    for j, t in enumerate(header):
        shape.table.cell(0, j).text = t
    for i, row in enumerate(rows, 1):
        for j, t in enumerate(row):
            shape.table.cell(i, j).text = t
    for i in range(len(rows) + 1):
        for j in range(len(header)):
            for p in shape.table.cell(i, j).text_frame.paragraphs:
                p.font.size = Pt(size)
    return s


# --------------------------------------------------------------------------------------------------------------
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Belief-based manipulation: five methods, one Bayes filter"
s.placeholders[1].text = ("How RSS20 (Wirnshofer), WAFR08 (Hsiao et al.), Audit-Before-You-Commit, GenAIF and "
                          "Quasi-static belief dynamics differ in what they believe, how they update it and how they act on it.\n"
                          "Plus: where skrl's PrivilegedPPO / BeliefPPO sit, the MuJoCo plan, and the encoder interface.")

equations("Common ground: the POMDP belief filter", [
    ("Belief = posterior over hidden state given history", r"b_t(s)=p(s_t=s\mid h_t),\qquad h_t=(o_0,a_0,\dots,a_{t-1},o_t)"),
    ("Prediction (transition model T)", r"b^-_{t+1}(s')=\int T(s'\mid s,a_t)\,b_t(s)\,ds"),
    ("Correction (observation model O)", r"b_{t+1}(s')=\frac{O(o_{t+1}\mid s',a_t)\;b^-_{t+1}(s')}{p(o_{t+1}\mid b_t,a_t)}"),
    ("Entropy (info-gathering criteria use it)", r"H(b)=-\sum_s b(s)\log b(s)"),
    ("Expected information gain of action a", r"\mathrm{IG}(a,b)=H(b)-\mathbb{E}_{o\sim p(o\mid b,a)}\big[H(b'_{a,o})\big]=I(s;o\mid b,a)"),
], subtitle="Every method below is this filter with a different choice of s, representation of b, T, O, and decision rule.",
    note="The methods differ on four axes: (1) what s is (object pose vs physical parameters vs discrete context vs "
         "learned latent), (2) how b is stored, (3) where T and O come from, (4) how b picks the action.")

equations("Same filter, different T and O", [
    ("RSS20: T = simulate each particle, O = robot feedback", r"q^{o,i}_{+}=\mathrm{sim}(q^{o,i},q^r,\kappa),\quad w^i\leftarrow w^i\,\mathcal{N}\!\big(\tilde q^r;\,\hat q^{r,i},\Sigma\big)"),
    ("WAFR08: T ≈ identity (light touch), O = offline table", r"b'(w)\propto \mathcal{N}\!\big(\langle\phi,c\rangle;\,\Omega_\tau(w,e),\Sigma\big)\,b(w)"),
    ("Audit: static parameters, T = identity, O = analytic f", r"b_k(\theta)\propto p_0(\theta)\prod_{j=1}^{k}\exp\!\Big(-\tfrac{\|y_j-f(u_j,\theta)\|^2}{2\sigma^2}\Big)"),
    ("GenAIF: T and O both learned (diffusion / CVAE)", r"b_{t+1}(z')\propto \ell_\theta(o_{t+1}\mid h_t,a_t,z')\sum_z M_{z'z}\,b_t(z)"),
    ("Quasi-static: no O at all (open loop), only T", r"b_+=\int p(q^o_+\mid q^o,u)\,b\,dq^o,\quad \mathrm{V}_+=\mathrm{V}_b[f(q^o,u)]+\mathbb{E}_b[\eta]\,\mathrm{V}_w"),
    ("skrl PrivilegedPPO: no explicit b, regressed point latent", r"\hat z_t=\hat E_\xi(o_{t-H+1:t},a_{t-H:t-1})\approx \mathbb{E}[z_t\mid h_t],\quad z_t=E_\psi(x^{\mathrm{priv}}_t)"),
], subtitle="The single most useful slide: which term of the Bayes filter each method keeps, learns, tabulates or drops.", cap_w=3.6)

table("At a glance", ["Method", "Belief over", "Representation", "Obs. model", "Action selection", "Information gathering",
                      "Commit / stop"], [
    ["RSS20 Wirnshofer (repo baseline)", "object pose (x, y[, yaw])", "300 particles → 2-comp. GMM", "sim rollout per particle vs robot feedback",
     "DQN picks a controller from a bank built around the GMM modes (2 Hz)", "implicit (learned Q, r = −1)", "goal reached"],
    ["WAFR08 Hsiao, Lozano-Pérez, Kaelbling", "object pose (x, y, θ)", "uniform grid, multinomial", "offline table Ω_τ(w,e) + Gaussian",
     "world-relative trajectory at most-likely state w*(b)", "explicit: argmax IG − c·FP", "b ∈ belief pre-image of goal (P ≥ 1−δ)"],
    ["Audit Before You Commit (Abouagour & Min 2026)", "task params θ (hole x, clearance; μ, m, k_d)", "1024 particles + retained set C_{k,ν}",
     "analytic forward model f(u, θ)", "probes, then minimax over C_{k,ν}", "fixed probe budget / cover schedule", "single irreversible commit; audit coverage + optimism"],
    ["GenAIF (Li, Jafari, Matta 2026)", "hidden dynamics context (tilt)", "discrete posterior, Markov switching M",
     "learned generative trajectory model", "MPC: min expected free energy over sampled U", "−λ·I(Ξ; Y) inside EFE", "none: replan every step"],
    ["Quasi-static belief dyn. (Jankowski et al. IJRR 2025)", "object pose q^o", "particles → variance only", "none (open loop)",
     "CMA-ES over via-points (BS-VP-STO)", "none: shrink variance by contact funnelling", "open-loop plan, γ_k ≤ 1"],
    ["skrl PrivilegedPPO (this fork)", "privileged params / state x^priv", "learned point latent ẑ", "implicit (regression to z)",
     "PPO policy on (o, ẑ)", "implicit", "—"],
], size=11, widths=[1.9, 1.7, 1.75, 1.8, 2.15, 1.75, 1.68])

# ---------------------------------------------------------------- RSS20
equations("RSS20 (Wirnshofer et al.) — the Belief-based-pushing baseline", [
    ("Hidden: object pose; robot joints observed with noise", r"x_t=(q^r_t,\dot q^r_t,q^o_t,\dot q^o_t),\qquad o_t=\tilde q^r_t=q^r_t+\epsilon"),
    ("Particle belief over object pose", r"b_t=\{(q^{o,i}_t,\,w^i_t)\}_{i=1}^{N_P},\qquad N_P=300"),
    ("Update: simulate every hypothesis under the same controller, compare robot feedback",
     r"w^i_{t+1}\propto w^i_t\;\mathcal{N}\!\big(\tilde q^r_{t+1};\;\hat q^{r,i}_{t+1}(q^{o,i}_t,\kappa_t),\;\Sigma\big)"),
    ("Compress (weighted EM, diagonal)", r"b_t\approx\sum_{g=1}^{N_G}\pi_g\,\mathcal{N}(\mu_g,\mathrm{diag}\,\sigma_g^2),\quad \tilde b_t=[\pi\,|\,\mu\,|\,\log\sigma^2]"),
    ("Controller bank relative to each mode", r"\mathcal{K}(\tilde b_t)=\{\kappa_{g,j}:\ \text{target}=\mu_g\oplus\delta_j\},\ g\le N_G,\ j\le N_\delta"),
    ("Policy: Dueling Double DQN on the compressed belief", r"\kappa_t=\arg\max_{k}Q_\phi(\tilde b_t,q^r_t,k),\quad r=-1,\ \gamma=0.98"),
], subtitle="Belief over WHERE the object is. Information gathering is never scored explicitly: the Q-function learns when probing pays.")

# ---------------------------------------------------------------- WAFR08
equations("WAFR08 — Robust belief-based execution (Hsiao, Lozano-Pérez, Kaelbling)", [
    ("Grid belief over object pose", r"w=(x,y,\theta)\in W\ \text{(cells)},\qquad b\in\Delta(W)"),
    ("Discrete Bayes filter", r"UB(b,a,o)(s')=\frac{\Pr(o\mid s',a)\sum_s\Pr(s'\mid s,a)\,b(s)}{\Pr(o\mid b,a)}"),
    ("Obs. model from offline geometric sim of WRT τ run for estimate e in true world w",
     r"P(\phi,c\mid w,\tau(e))=\mathcal{N}\big(\Omega_\tau(w,e),\Sigma\big)"),
    ("Act at the most likely state", r"w^*(b)=\arg\max_w b(w)"),
    ("Goal on beliefs, and the belief pre-image", r"G_\delta:\ \sum_w b(w)\,\mathbb{I}[G(\phi,w)]>1-\delta,\qquad \mathrm{BPre}(G_\delta,\tau(e))=\{b:\ P_b(G(R(\tau(e),w)))>1-\delta\}"),
    ("Failure probability (undetected contact)", r"FP(\tau(e),b)=P_b\big(R(\tau(e),w)=\mathit{fail}\big)"),
    ("Information gain of a WRT (first contact only)", r"IG(\tau,b)=H(b)-\sum_j b(w_j)\,H\Big(UB\big(b,\tau(w^*(b)),\Omega_\tau(w_j,w^*(b))\big)\Big)"),
    ("Mode switch", r"\begin{cases}\tau^* \text{ s.t. } b\in\mathrm{BPre}(G_\delta,\tau^*(w^*(b))) & \text{goal mode}\\ \tau^*=\arg\max_\tau IG(\tau,b)-c\,FP(\tau,b) & \text{info mode}\end{cases}"),
], subtitle="Belief over WHERE the object is. Explicit, one-step information gain; commits when the belief is inside the goal's pre-image.",
    size=20, cap_w=3.5)

# ---------------------------------------------------------------- Audit
equations("Audit Before You Commit (Abouagour & Min, arXiv 2609.30608)", [
    ("Hidden task parameters (insertion example; pushing: μ, m, k_d)", r"\theta^\star\sim P_{\mathrm{env}},\quad \theta=(x_{\mathrm{hole}},c),\quad b_0=p_0"),
    ("Particle belief from k probes (static θ: no prediction step)", r"b_k(\theta)\propto p_0(\theta)\prod_{j=1}^{k}\tilde p_f(y_j\mid u_j,\theta),\quad \tilde p_f\propto\exp\!\Big(-\frac{\|y_j-f(u_j,\theta)\|^2}{2\sigma^2}\Big)"),
    ("Retained set: shortest top-weight prefix with mass ν (=0.8)", r"C_{k,\nu}=\{\theta_{(1)},\dots,\theta_{(m)}\},\quad m=\min\{m:\textstyle\sum_{i\le m}w_{(i)}\ge\nu\}"),
    ("Commit once: scenario minimax (binary: failure fraction)", r"a^C_{k,\nu}=\arg\min_a\max_{\theta\in C_{k,\nu}}\mathrm{miss}(a,\theta),\qquad \hat R(a)=\frac{1}{|C|}\sum_{\theta\in C}\mathbb{I}[\mathrm{fail}(a,\theta)]"),
    ("Probe budget trade-off", r"J_\lambda(k)=\mathrm{miss}(k)+\lambda k"),
    ("Audit 1: coverage of the decision coordinate", r"\Pr\big[z^\star\in I(C_{k,\nu})\big]\overset{?}{\ge}\nu\quad(\text{support escape: } z^\star\notin\mathrm{hull})"),
    ("Audit 2: failure-model optimism", r"e_k=R_k-\hat R_k\quad(>0\Rightarrow\text{model optimistic})"),
    ("Limit of any history-based commit gate", r"\big|\Pr_{\theta^\star}[g(H)=1]-\Pr_{\tilde\theta}[g(H)=1]\big|\le\mathrm{TV}\big(P^H_{\theta^\star},P^H_{\tilde\theta}\big)"),
], subtitle="Belief over WHAT the physics/geometry is. Not a planner: it asks whether the belief can be trusted before the one irreversible action.",
    size=20, cap_w=3.6)

# ---------------------------------------------------------------- GenAIF
equations("GenAIF — generative models for active inference (Li et al., arXiv 2610.05692)", [
    ("One generative model gives proposals, dynamics and likelihood",
     r"p_\theta(U,S_+,Y\mid c_t)=\prod_{k=t}^{t+H-1}\mu_\theta(a_k\mid c_k)\,T_\theta(s_{k+1}\mid s_k,a_k)\,O_\theta(o_{k+1}\mid s_{k+1},a_k)"),
    ("Belief over hidden context z (benchmark: lateral tilt)", r"b_{t+1}(z')\propto\ell_\theta(o_{t+1}\mid h_t,a_t,z')\sum_z M_{z'z}\,b_t(z),\quad z\in\{-15^\circ,0^\circ,15^\circ\}"),
    ("Update as variational free-energy minimisation", r"F_{t+1}[q]=D_{\mathrm{KL}}\big(q\,\|\,b^-_{t+1}\big)-\mathbb{E}_q\big[\log O_\theta(o_{t+1}\mid s',a_t)\big]"),
    ("Candidate rollout distribution under do(U)", r"Q^U_t(\mathbf{S},Y)=b_t(s_t)\prod_k T_\theta(s_{k+1}\mid s_k,a_k)\,O_\theta(o_{k+1}\mid s_{k+1},a_k)"),
    ("Expected free energy: task cost minus info gain", r"G_t(U)=\underbrace{\mathbb{E}_{Q^U_t}\big[C(U,Y;g)\big]}_{\text{pragmatic}}-\lambda\underbrace{I_{Q^U_t}(\Xi;Y)}_{\text{epistemic}}"),
    ("MPC: sample U from μ_θ, keep safe ones, take min", r"U^*=\arg\min_{U\sim\mu_\theta,\ \mathrm{risk}(U)\le\rho}G_t(U),\quad a_t=U^*_0,\ \text{then update } b,\ \text{replan}"),
], subtitle="Belief over WHICH dynamics regime holds. Info gain is a term in the objective, traded off continuously against task cost.",
    size=20, cap_w=3.4,
    note="Note: the paper's hidden variable is surface tilt (9 conditions incl. obstacle position), not friction/mass. "
         "μ, m, damping are a natural extension in our MuJoCo env.")

# ---------------------------------------------------------------- Quasi-static
equations("Quasi-static belief dynamics for robust pushing (Jankowski et al., IJRR 2025)", [
    ("Quasi-static contact: object moves the least to stay non-penetrating",
     r"q^o_+=q^o+\delta q^o,\quad \delta q^o=\arg\min_{\delta}\ \delta^\top M(q^o)\,\delta\ \ \text{s.t.}\ d(u,q^o+\delta)\ge0"),
    ("Noise only while in contact", r"q^o_+=f(q^o,u)+\eta\,w,\quad \eta=\mathbb{I}[d(u,q^o)\le0],\quad \mathbb{E}[w]=0"),
    ("Variance propagation without sampling w", r"\mathrm{V}_+=\mathrm{V}_b\big[f(q^o,u)\big]+\mathbb{E}_b[\eta]\,\mathrm{V}_w"),
    ("Particle estimate (nominal f only)", r"\hat{\mathrm{V}}_+=\frac{1}{N_p}\sum_i\big(f(q^{o,i},u)-\bar\mu\big)^{2}+\Big(\frac{1}{N_p}\sum_i\eta^i\Big)\mathrm{V}_w"),
    ("Variance gain per step; robustness constraint", r"\gamma_k=\frac{\hat{\mathrm{V}}_{\tilde b_{k+1}}+\hat{\mathbb{E}}_{\tilde b_k}[\eta]\,\mathrm{V}_w}{\hat{\mathrm{V}}_{\tilde b_k}+\mathrm{V}_w}\le1\quad\forall k"),
    ("Cost for CMA-ES over via-points (BS-VP-STO)", r"c=\lambda_c\prod_{k}\exp\!\Big(-\frac{1-\gamma_k}{K-1}\Big)+\big\|\mathbb{E}[q^o_K]-q^o_{\mathrm{des}}\big\|^2,\quad \lambda_c\in\{1,10^3\}"),
], subtitle="Belief over WHERE the object is, but never corrected by sensing: robustness comes from choosing contacts that funnel (contract) the belief.",
    size=20, cap_w=3.4)

# ---------------------------------------------------------------- skrl
equations("Where this skrl fork sits: PrivilegedPPO / BeliefPPO", [
    ("Teacher latent from privileged values (sim only)", r"z_t=E_\psi(x^{\mathrm{priv}}_t)"),
    ("History estimator (MLP / Conv1d / LSTM / Transformer)", r"\hat z_t=\hat E_\xi\big(o_{t-H+1:t},a_{t-H:t-1}\big)"),
    ("Policy input, masked by curriculum", r"\pi\big(a_t\mid o_t,\ m_t z_t+(1-m_t)\hat z_t\big),\quad m_t\sim\mathrm{Bern}(1-p_{\mathrm{mask}})"),
    ("Estimator / sysid losses", r"\mathcal{L}=\big\|\hat z_t-\mathrm{sg}(z_t)\big\|^2+\beta\,\big\|\hat x^{\mathrm{priv}}_t-x^{\mathrm{priv}}_t\big\|^2"),
    ("What the latent is, in belief terms", r"\hat z_t\approx\mathbb{E}[z_t\mid h_t]\quad\text{(a point estimate: no }\mathrm{Var}[z_t\mid h_t]\text{, no info-gain term)}"),
    ("BeliefPPO (explicit): particle filter + GMM as obs.", r"\pi\big(a_t\mid o_t,\tilde b_t\big),\quad \tilde b_t=[\pi_g\,|\,\mu_g\,|\,\log\sigma_g^2]_{g\le N_G}"),
], subtitle="Amortised belief: fast and scalable (1000s of envs), but uncertainty and information value are implicit in the RL value function.")

bullets("How they differ, in belief terms", [
    "WHAT is uncertain",
    "  pose of the object: RSS20, WAFR08, Quasi-static;  physical / task parameters: Audit (θ), GenAIF (tilt), skrl (x^priv)",
    "HOW uncertainty is stored",
    "  full posterior: WAFR08 grid, Audit particles;  compressed: RSS20 2-GMM;  2nd moment only: Quasi-static;  point: skrl ẑ",
    "WHERE the likelihood comes from",
    "  offline table (WAFR08), analytic model (Audit), per-particle simulation (RSS20), learned generative model (GenAIF), none (Quasi-static)",
    "HOW information is valued",
    "  entropy drop IG(τ,b) (WAFR08) · mutual information I(Ξ;Y) inside EFE (GenAIF) · fixed probe budget + coverage audit (Audit)",
    "  · variance gain γ_k ≤ 1 (Quasi-static) · learned implicitly by Q / V (RSS20, skrl)",
    "WHEN to act on it",
    "  most-likely state + pre-image test (WAFR08) · worst case over retained set (Audit) · expectation (GenAIF) · every step (RSS20, skrl)",
], size=17)

bullets("Proposed MuJoCo implementation in skrl (to agree before coding)", [
    "One env: planar cube push (port of Belief-based-pushing mujoco_repro), hidden pose + μ, m, k_d, batched, debug video",
    "Shared belief core: skrl.belief ParticleFilter (exists) + grid filter + retained-set / coverage utilities",
    "Methods as pluggable deciders on the same filter:",
    "  wafr08:  WRT library, offline Ω-table, IG − c·FP vs. pre-image test",
    "  audit:   probe-then-commit, scenario minimax, coverage + optimism report",
    "  genaif:  CVAE trajectory model (fast) → likelihood + EIG, EFE-MPC",
    "  quasistatic: nominal-particle variance propagation, γ_k constraint, CMA-ES over via-points",
    "  rss20:   existing BeliefDQN on the 2-GMM (baseline)",
    "Benchmarks: success / miss vs. probes, belief coverage, wall-clock; sweeps on dl",
], size=17)

equations("Encoders: pluggable observation encoders (DINO, CLIP, PointNet, FiLM, any)", [
    ("Dict observation, one encoder per key", r"o_t=\{\text{rgb},\ \text{cloud},\ \text{text},\ \text{proprio},\ \tilde b_t\},\quad h^k=\mathrm{Enc}_k(o^k_t)"),
    ("Frozen pretrained (DINOv2 / CLIP image)", r"h^{\mathrm{rgb}}=\mathrm{sg}\big(\mathrm{ViT}(\mathrm{norm}_{\mathrm{ImageNet}}(\text{rgb}))\big)\in\mathbb{R}^{d}"),
    ("PointNet (permutation invariant)", r"h^{\mathrm{cloud}}=\max_{i\le N}\ \mathrm{MLP}(p_i)"),
    ("CLIP text as condition", r"c=\mathrm{CLIP}_{\mathrm{text}}(\text{instruction})"),
    ("FiLM fuses condition into features", r"\mathrm{FiLM}(h\mid c)=\gamma(c)\odot h+\beta(c)"),
    ("Per-key preprocessors (scaler only on low-dim keys)", r"\tilde o^k=\begin{cases}(o^k-\mu_k)/\sigma_k & k\in\{\text{proprio},\tilde b\}\\ o^k & \text{pixels, points, tokens}\end{cases}"),
], subtitle="Today skrl flattens dict observations and RunningStandardScaler would normalise pixels and token ids; PrivilegedPPO requires no preprocessor.",
    note="YAML: models.policy.encoders: {rgb: {type: dino, frozen: true}, text: {type: clip_text}, cloud: {type: pointnet}}, "
         "fusion: {type: film, condition: text}")

bullets("Sources", [
    "K. Hsiao, T. Lozano-Pérez, L. P. Kaelbling. Robust Belief-Based Execution of Manipulation Programs. WAFR 2008. lis.csail.mit.edu/pubs/tlp/wafr08.pdf",
    "L. Wirnshofer, P. Schmitt, G. v. Wichert, W. Burgard. Controlling Contact-Rich Manipulation Under Partial Observability. RSS 2020.",
    "M. Abouagour, B.-C. Min. Audit Before You Commit: Locating Belief Failures in Active Identification for One-Shot Manipulation. arXiv 2609.30608.",
    "Y. Li, M. A. Jafari, A. Matta. Benchmarking Generative Trajectory Models for Active-Inference Control. arXiv 2610.05692.",
    "J. Jankowski et al. Robust Pushing: Exploiting Quasi-static Belief Dynamics and Contact-informed Optimization. IJRR 2025, arXiv 2404.02795.",
    "github.com/ebasatemesgen/Belief-based-pushing (mujoco_repro = RSS20 MuJoCo reproduction track).",
], size=14)

out = os.path.join(HERE, "Belief_methods_equations.pptx")
prs.save(out)
print(out, len(prs.slides), "slides")
