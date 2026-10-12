"""Build docs/belief_methods/Cube_push_belief_methods.pptx: setup definitions (observation, action and reward spaces)
and WAFR08, Audit, GenAIF, quasi-static belief dynamics and RSS20 on the same MuJoCo cube push, with belief renders
(uploaded to Google Slides afterwards; the mp4s are embedded there from Drive).

    MUJOCO_GL=egl .venv/bin/python examples/contact_rich_uncertainty/cube_push_belief.py --method <m> --seed {0,1,2}
    .venv/bin/python docs/belief_methods/build_cube_push_methods_slides.py
"""
import json
import os
import subprocess

import imageio_ffmpeg
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

plt.rcParams.update({"text.usetex": True, "text.latex.preamble": r"\usepackage{amsmath,amssymb,bm}"})

HERE = os.path.dirname(os.path.abspath(__file__))
RUNS = os.path.join(HERE, "..", "..", "runs")
IMG = os.path.join(HERE, "cube_push_methods")
os.makedirs(IMG, exist_ok=True)
SEEDS = (0, 1, 2)
FRACTIONS = (0.02, 0.3, 0.6, 0.995)
METHODS = {"wafr08": "WAFR08", "audit": "Audit", "genaif": "GenAIF", "quasistatic": "Quasi-static", "rss20": "RSS20"}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
GREY = RGBColor(0x55, 0x55, 0x55)


def title_only(title, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    t = s.shapes.title
    t.text = title
    t.text_frame.paragraphs[0].font.size = Pt(28)
    t.left, t.top, t.width, t.height = Inches(0.4), Inches(0.15), W - Inches(0.8), Inches(0.8)
    if subtitle:
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.9), W - Inches(1), Inches(0.4))
        tb.text_frame.word_wrap = True
        tb.text_frame.text = subtitle
        tb.text_frame.paragraphs[0].font.size = Pt(14)
        tb.text_frame.paragraphs[0].font.color.rgb = GREY
    return s


def text(s, x, y, w, h, lines, size=15):
    tf = s.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        level = (len(line) - len(line.lstrip(" "))) // 2
        p.text, p.level = line.strip(), level
        p.font.size = Pt(size - 2 * level)
    return tf


def latex(name, tex, size=20):
    path = os.path.join(IMG, f"{name}.png")
    fig = plt.figure()
    fig.text(0, 0, tex, fontsize=size)
    fig.savefig(path, bbox_inches="tight", pad_inches=0.05, dpi=200, transparent=True)
    plt.close(fig)
    return path


def table(s, x, y, w, rows, size=12, col_widths=None):
    shape = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(0.3 * len(rows)))
    tbl = shape.table
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = tbl.cell(i, j)
            c.text = str(cell)
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(size)
                p.font.bold = i == 0
    if col_widths:
        for j, cw in enumerate(col_widths):
            tbl.columns[j].width = Inches(cw)
    return tbl


def keyframes(method, seed):
    video = os.path.join(RUNS, f"cube_push_{method}", "videos", f"{method}_seed{seed}.mp4")
    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    n = imageio_ffmpeg.count_frames_and_secs(video)[1]
    paths = []
    for k, f in enumerate(FRACTIONS):
        path = os.path.join(IMG, f"{method}_seed{seed}_{k}.jpg")
        subprocess.run([ffmpeg, "-y", "-loglevel", "error", "-ss", f"{f * n:.2f}", "-i", video, "-frames:v", "1",
                        "-q:v", "3", path], check=True)
        paths.append(path)
    return paths


def load(method):
    out = {}
    for seed in SEEDS:
        path = os.path.join(RUNS, f"cube_push_{method}", f"seed{seed}.json")
        if os.path.exists(path):
            out[seed] = json.load(open(path))
    return out


def distance_plot(results):
    path = os.path.join(IMG, "distance.png")
    plt.rcParams["text.usetex"] = False
    fig, axes = plt.subplots(1, len(METHODS), figsize=(13, 3.3), dpi=130, sharey=True)
    for a, (m, name) in zip(axes, METHODS.items()):
        for seed, r in results[m].items():
            log = r["log"]
            a.plot(range(1, len(log) + 1), [100 * l["goal_dist"] for l in log], ".-", ms=3, label=f"seed {seed}")
        a.axhline(2.0, color="g", ls="--", lw=1)
        a.set_title(name)
        a.set_xlabel("controller #")
        a.grid(alpha=0.3)
        a.legend(fontsize=7)
    axes[0].set_ylabel("true distance to goal [cm]")
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    plt.rcParams["text.usetex"] = True
    return path


def summary(r):
    kinds = [l["wrt"].split()[0] for l in r["log"]]
    pushes = sum(k in ("goal", "qs") for k in kinds)
    return (f"{'yes' if r['success'] else 'no'}", f"{r['goal_dist_cm']:.1f} cm / {r['goal_dyaw_deg']:.0f}°",
            f"{len(kinds)} ({kinds.count('touch')} / {kinds.count('probe')} / {pushes})", f"{r['sim_time_s']:.0f} s",
            f"{r['wall_s']:.0f} s")


results = {m: load(m) for m in METHODS}

# ---------------------------------------------------------------- title + summary
s = title_only("Cube push: five belief-space methods on one MuJoCo env",
               "Hidden cube pose + friction μ, mass m, damping k_d; the only command is the wanted final cube pose; "
               "same WRT library, same simulator, same debug video for every method")
rows = [["method", "seed", "success", "final error", "controllers (touch / probe / push)", "sim time", "wall"]]
for m, name in METHODS.items():
    for seed, r in results[m].items():
        rows.append([name, seed, *summary(r)])
table(s, Inches(0.5), Inches(1.5), Inches(12.3), rows, size=11, col_widths=[1.6, 0.7, 1.0, 1.9, 3.6, 1.4, 1.4])

# ---------------------------------------------------------------- definitions: env
s = title_only("Setup definitions: the environment (CubePushEnv, gymnasium)",
               "skrl/envs/cube_push/env.py; MuJoCo 2 ms physics step, 10 substeps → 50 Hz control")
text(s, Inches(0.5), Inches(1.4), Inches(12.4), Inches(6.0), [
    "Hidden state s = (x, y, ψ, μ, m, k_d)",
    "  initial cube pose x ∈ [0, 0.12] m, y ∈ [−0.10, 0.10] m, ψ any (square: ψ mod 90°); 6 cm cube",
    "  friction μ ∈ [0.3, 0.8], mass m ∈ [0.2, 0.6] kg, damping k_d ∈ [0, 3] N s/m (k_d h² on yaw); all uniform per episode",
    "Command: goal pose (x_g, y_g, ψ_g), x_g ∈ [0.18, 0.30] m, y_g ∈ [−0.12, 0.12] m, ψ_g any",
    "Observation space: Box(10), float32",
    "  pusher position p ∈ R³ (encoder noise σ = 0.5 mm), pusher actuator force F ∈ R³ (σ = 0.05 N)",
    "  goal (x_g, y_g, cos 4ψ_g, sin 4ψ_g); no cube pose, no physics (info carries them for evaluation and rendering only)",
    "Action space: Box(3), pusher position target (x, y, z)",
    "  x ∈ [−0.45, 0.45], y ∈ [−0.35, 0.35], z ∈ [0.033, 0.2] m; change per step clipped to 0.5 m/s × 20 ms = 1 cm",
    "Reward: r_t = −‖p_cube − p_g‖ − 0.05 |ψ − ψ_g|₉₀ − 10 · 1[cube off table]  (m, rad; r ∈ [≈ −11.2, 0])",
    "Termination: cube off the table; truncation: 4000 steps (80 s); info['success'] = within 2 cm and 10°",
    "The planners below do not use r: they act through world-relative trajectories (next slide) and are scored by "
    "the true final error and info['success']",
], size=16)

# ---------------------------------------------------------------- definitions: controller level
s = title_only("Setup definitions: controller level used by the planners",
               "skrl/envs/cube_push/wafr08.py (WRT library, features, grid belief); every method executes these "
               "through the env above")
text(s, Inches(0.5), Inches(1.4), Inches(6.3), Inches(6.0), [
    "Actions: world-relative trajectories (WRT) at an estimated cube state w",
    "  guarded touch (face f ∈ 0..3, offset ∈ {−½, 0, ½} h): stop at contact",
    "  force probe (face f): push 4 cm at 2 speeds, measure forces",
    "  goal push (face, offset ∈ {−0.6, 0, 0.6} h, v ∈ {0.05, 0.25} m/s, L ∈ {1, 2, 4, 7, 10} cm)",
    "  quasi-static push (approach angle, offset, length; CMA-ES decision variables)",
    "Observations: features of the sensor trace",
    "  touch (c, travel); goal push (c, travel, F̄₁); probe (c, travel, F̄₁, F̄₂, F_peak)",
    "  c ∈ {0 no contact, 1 contact, 2 pusher landed on the cube}",
    "  F̄ ≈ μ m g + k_d v (steady push), F_peak ≈ m a (start)",
    "Stop / success: P_b(within 2 cm, 10°) ≥ 1 − δ (δ = 0.2), budget 40 controllers",
], size=15)
text(s, Inches(6.9), Inches(1.4), Inches(6.0), Inches(6.0), [
    "Beliefs",
    "  grid 36 × 35 × 9 × 3 × 3 × 3 over (x, y, ψ, μ, m, k_d), 1.5 cm / 10° cells: WAFR08, Audit, GenAIF, quasi-static",
    "  128 particles (physics on 5-level lattices): RSS20",
    "Costs",
    "  goal cost d / 2 cm + |Δψ|₉₀ / 10° (WAFR08, Audit, GenAIF)",
    "  quasi-static: belief contraction + mean-pose error (CMA-ES)",
    "RSS20 controller MDP (learned policy)",
    "  observation Box(34): b̃ = 2-component GMM [π (2) | μ (2 × 7) | log σ² (2 × 7)] + goal (4)",
    "  action Discrete(16): per GMM mode 4 touches, 1 probe, 3 goal pushes",
    "  reward −1 per controller + γΦ(s′) − Φ(s), Φ = −d / 2 cm (training only), −10 off table",
    "  γ = 0.98; terminal at true success or off table; ≤ 30 controllers",
], size=15)

# ---------------------------------------------------------------- per method
RULES = {
    "audit": ("Audit Before You Commit (Abouagour & Min 2026)", [
        r"$C_{k,\nu}$ = shortest top-weight prefix of grid cells holding mass $\nu = 0.8$ (retained set)",
        r"probe: fixed touch schedule while $C$ is not tight (weighted RMS $> 2$ cm or $> 15^\circ$), $\leq 5$ per phase",
        r"commit: $\tau^* = \arg\min_\tau \max_{w \in C} J(T_\tau(w))$ over the 12 best goal pushes at $w^*$",
        r"$\hat R$ = fraction of $C$ whose outcome misses the nominal outcome; audit with the truth: "
        r"coverage $1[s \in C]$, $e = R - \hat R$",
    ], ["Paper commits once; one push cannot reach the goal here, so probe → commit repeats.",
        "Audit needs the true state: evaluation only (runner sets planner.true_state)."]),
    "genaif": ("Generative active inference (GenAIF)", [
        r"learned $p_\theta(o, \Delta w \mid \tau, w, \mu, m, k_d)$: MLP 12 → 3 × 256 → contact logit + 7 Gaussians",
        r"trained on 45 830 MuJoCo rollouts (12 WRT profiles): contact acc. 0.98, travel MAE 2.5 mm",
        r"belief update with $p_\theta$, $\sigma_{\rm eff}^2 = \sigma_{\rm sensor}^2 + \sigma_\theta^2$, "
        r"Markov switching mix on physics",
        r"$\tau^* = \arg\min_\tau\ \mathbb{E}_b[J(T_\tau)] - \lambda\, IG(\tau)$, \ $\lambda = 2$, risk filter "
        r"$P(\mathrm{fail}) \leq 0.2$",
    ], ["Model replaces the simulator in the belief update and the plan (no MuJoCo at decision time for "
        "predictions).", "No stop rule fires in practice: belief stays broad (learned σ), budget ends the run."]),
    "quasistatic": ("Quasi-static belief dynamics (Jankowski et al., IJRR 2025)", [
        r"localize with touches (IG $-$ c FP) while $H(\mathrm{pose}) > 3$ nats",
        r"plan $K = 3$ slow pushes by CMA-ES on 32 nominal particles:",
        r"$c = \prod_k e^{-(1-\gamma_k)/(K-1)} + \lambda_c \sum_k [\gamma_k - 1]_+ + "
        r"\|\bar p_K - p_g\|^2/\mathrm{tol}_p^2 + \Delta\bar\psi_K^2/\mathrm{tol}_\psi^2$",
        r"$\gamma_k = (\hat V_{k+1} + \hat E_k[\eta] V_w) / (\hat V_k + V_w)$; run open loop "
        r"(belief follows $T$ only), then re-localize",
    ], ["Paper's 10³ switch on γ > 1 made CMA-ES choose pushes that never touch the cube (unknown μ, m spread "
        "the belief): linear penalty λ_c = 10 instead.", "Paper assumes a roughly known start; the any-yaw "
        "12 × 20 cm prior needs the localization phase."]),
    "rss20": ("RSS20: belief-space controller selection (Wirnshofer et al. 2020)", [
        r"particle belief (128), each particle simulated in MuJoCo under the executed controller, "
        r"weighted by $P(o \mid s)$, resampled with jitter",
        r"$\tilde b$ = 2-component diagonal GMM of standardized $(x, y, \cos 4\psi, \sin 4\psi, \mu, m, k_d)$, heaviest first",
        r"bank of 16 controllers around the 2 modes; policy $\arg\max_a Q_\theta(\tilde b, g, a)$",
        r"$Q_\theta$: dueling double DQN (ResNet 4 × 64), 48 parallel envs on dl, 72k transitions",
    ], ["Stops on the task-completion signal (true success), as in the paper; the other methods stop on belief.",
        "Training reward is shaped by the true distance (privileged, training only)."]),
}
for m in ("audit", "genaif", "quasistatic", "rss20"):
    name = METHODS[m]
    title, rule, notes = RULES[m]
    s = title_only(f"{name}: method on the cube push", title)
    s.shapes.add_picture(latex(f"{m}_rule", "\n".join(rule), size=18), Inches(0.6), Inches(1.5), width=Inches(12.0))
    seeds = results[m]
    text(s, Inches(0.6), Inches(5.0), Inches(12.2), Inches(2.4), [
        *notes,
        "Results: " + "; ".join(f"seed {k}: {'success' if r['success'] else 'fail'}, {r['goal_dist_cm']:.1f} cm "
                                f"{r['goal_dyaw_deg']:.0f}°, {len(r['log'])} controllers" for k, r in seeds.items()),
    ], size=14)
    if m == "audit" and seeds:
        a = [r.get("audit", {}) for r in seeds.values()]
        text(s, Inches(0.6), Inches(6.6), Inches(12.2), Inches(0.8), ["Audit: " + "; ".join(
            f"seed {k}: {x.get('commits', 0)} commits, coverage {x.get('coverage') or 0:.2f}, outcome coverage "
            f"{x.get('outcome_coverage', 0):.2f}, miss {x.get('miss_rate', 0):.2f} vs R^ "
            f"{x.get('predicted_miss_rate', 0):.2f}" for k, x in zip(seeds, a))], size=12)
    if 1 in seeds:
        r = seeds[1]
        mu, mass, kd = r["physics"]
        s = title_only(f"{name}: belief over time (seed 1 keyframes)",
                       f"true μ {mu:.2f}, m {mass:.2f} kg, k_d {kd:.2f} | final error {r['goal_dist_cm']:.1f} cm, "
                       f"{r['goal_dyaw_deg']:.0f}° | stop: {r['stopped'][:70]}")
        for k, path in enumerate(keyframes(m, 1)):
            s.shapes.add_picture(path, Inches(0.35 + 6.35 * (k % 2)), Inches(1.35 + 3.05 * (k // 2)),
                                 height=Inches(3.0))
        text(s, Inches(4.15), Inches(1.35), Inches(2.4), Inches(5.9), [
            "Orange tiles: belief cells (opacity ∝ mass)",
            "Blue box: estimate w*",
            "Green: goal pose, 2 cm disc",
            "Yellow: current controller path",
            {"audit": "Purple dots: retained set C", "genaif": "Belief from the learned model",
             "quasistatic": "Cyan: planned mean poses of the open-loop plan",
             "rss20": "Magenta dots: particles; boxes: GMM modes"}[m],
        ], size=11)
    title_only(f"{name}: full renders (seeds 0, 1, 2)", f"Drive videos {m}_seed{{0,1,2}}.mp4 "
               f"(runs/cube_push_{m}/videos/)")

# ---------------------------------------------------------------- comparison
s = title_only("Comparison: true distance to the goal per controller",
               "green: 2 cm tolerance; WAFR08 and Audit stop on belief, RSS20 on the task signal, GenAIF and "
               "quasi-static mostly run to the 40-controller budget")
s.shapes.add_picture(distance_plot(results), Inches(0.2), Inches(1.5), width=Inches(12.9))
text(s, Inches(0.6), Inches(5.0), Inches(12.2), Inches(2.3), [
    "Same env, WRT library, simulator and render for all; 3 seeds each (seed = cube pose, physics and goal).",
    "Test implementations: hyper-parameters set by hand on seed 1, not tuned; 3 seeds do not support ranking claims.",
], size=15)

# ---------------------------------------------------------------- limits
s = title_only("What the test implementations show, and what they do not")
text(s, Inches(0.6), Inches(1.3), Inches(12.2), Inches(5.9), [
    "Shows",
    "  All five methods run end to end on the same partially observable MuJoCo cube push with one runner "
    "(cube_push_belief.py --method) and the same belief video",
    "  Audit's coverage check exposes when the grid belief is overconfident (truth outside C)",
    "Does not yet",
    "  Physics identification stays coarse in every method (push forces mostly measure μ·m)",
    "  GenAIF: learned σ keeps the belief broad, so its belief-based stop rarely fires",
    "  Quasi-static: open-loop plans with unknown μ, m expand the belief (γ > 1); many re-localizations",
    "  RSS20: one short DQN run (~1.7 h CPU), no hyper-parameter search",
    "Code: skrl/envs/cube_push/{wafr08,audit,genaif,quasistatic,rss20,render}.py, "
    "examples/contact_rich_uncertainty/{cube_push_belief,cube_push_rss20_train}.py",
    "Runs on zz-bw (planners) and dl (RSS20 training), CPU; videos in runs/cube_push_<method>/videos/",
], size=17)

out = os.path.join(HERE, "Cube_push_belief_methods.pptx")
prs.save(out)
print(out, len(prs.slides), "slides")
