"""Build docs/privileged_curriculum/History_encoder_sysid.pptx (uploaded to Google Slides afterwards).

    .venv/bin/python docs/privileged_curriculum/plot_gym.py runs/privileged_gym docs/privileged_curriculum/diagrams
    .venv/bin/python docs/privileged_curriculum/build_sysid_slides.py runs/privileged_gym
"""
import collections
import glob
import json
import os
import sys

import numpy as np
from pptx import Presentation
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
D = os.path.join(HERE, "diagrams")
root = sys.argv[1]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height


def title_only(title):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = title
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.4), Inches(0.2), W - Inches(0.8), Inches(0.9)
    return s


def bullets(title, items, size=20):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.4), Inches(0.2), W - Inches(0.8), Inches(1.0)
    body = s.placeholders[1]
    body.left, body.top, body.width, body.height = Inches(0.6), Inches(1.4), W - Inches(1.2), H - Inches(1.8)
    tf = body.text_frame
    tf.clear()
    for i, it in enumerate(items):
        level = 1 if it.startswith("  ") else 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.level = it.strip(), level
        p.font.size = Pt(size - 3 * level)
    return s


def image(title, png, note=None):
    s = title_only(title)
    from PIL import Image

    iw, ih = Image.open(png).size
    top = Inches(1.15)
    max_w, max_h = W - Inches(0.6), H - top - (Inches(0.6) if note else Inches(0.2))
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    s.shapes.add_picture(png, Emu((W - w) // 2), top, w, h)
    if note:
        tb = s.shapes.add_textbox(Inches(0.5), H - Inches(0.6), W - Inches(1), Inches(0.5))
        tb.text_frame.text = note
        tb.text_frame.paragraphs[0].font.size = Pt(14)
    return s


def table(title, header, rows, size=13):
    s = title_only(title)
    shape = s.shapes.add_table(len(rows) + 1, len(header), Inches(0.4), Inches(1.2), W - Inches(0.8), Inches(0.4))
    for j, text in enumerate(header):
        shape.table.cell(0, j).text = text
    for i, row in enumerate(rows, 1):
        for j, text in enumerate(row):
            shape.table.cell(i, j).text = text
    for cell in (shape.table.cell(i, j) for i in range(len(rows) + 1) for j in range(len(header))):
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(size)
    return s


runs = [json.load(open(p)) for p in sorted(glob.glob(os.path.join(root, "*.json")))]
ablation = [json.load(open(p)) for p in sorted(glob.glob(os.path.join(root, "ablation", "*.json")))]
by = collections.defaultdict(list)
for r in runs:
    by[r["task"], r["mode"]].append(r)


def mean_std(vals):
    return f"{np.mean(vals):.0f} ± {np.std(vals):.0f}"


s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "History encoder of privileged information = online system identification"
s.shapes.title.left, s.shapes.title.top = Inches(0.8), Inches(2.0)
s.shapes.title.width, s.shapes.title.height = W - Inches(1.6), Inches(1.5)
s.placeholders[1].left, s.placeholders[1].top = Inches(0.8), Inches(3.8)
s.placeholders[1].width, s.placeholders[1].height = W - Inches(1.6), Inches(1.8)
s.placeholders[1].text = (
    "PrivilegedPPO (ROA-style) with a privileged-masking curriculum: what the history estimator recovers\n"
    "skrl fork · branch feature/privileged-curriculum · 2026-10-11"
)

bullets("Problem", [
    "Sim-to-real: the policy needs quantities the real robot cannot measure",
    "  Hidden state (velocities from noisy positions) and dynamics (mass, friction, link length)",
    "Training in simulation can read them: privileged information e_t",
    "Deployment cannot: the policy must infer them online from what it has seen",
    "  Inferring dynamics from an observation (and action) history = online system identification",
    "Question: which privileged groups does a history encoder actually recover, and what helps it?",
])
bullets("RMA vs ROA vs this fork", [
    "RMA (Kumar et al. 2021): two phases",
    "  Phase 1: train policy on encoder latent z = μ(e_t). Phase 2: regress ẑ = φ(history) onto z, freeze policy",
    "ROA (Fu et al. 2022, Deep Whole-Body Control): one phase",
    "  Train z and ẑ together with ||sg(z) − ẑ||² + λ||z − sg(ẑ)||²; the policy sees z or ẑ",
    "This fork: ROA + a per-episode mask curriculum",
    "  Mask flag drawn per episode: masked → actor uses ẑ, else z. Critic always sees e_t (asymmetric)",
    "  Time-based softmax schedule, or reward-gated: mask p grows one step only while the return holds",
    "  Probes: held-out R² of a small MLP from ẑ (and z) to each privileged group (diagnostic only)",
])
image("Pipeline: the history estimator is the online system-ID module", os.path.join(D, "02_history_sysid.png"),
      "Deploy: mask = 1, the estimator φ runs on the real robot's history only")
bullets("Losses and the curriculum", [
    "PPO loss on the actor that acts on z (visible episodes) or ẑ (masked episodes)",
    "Estimator: ||sg(z) − ẑ||²   (ẑ learns to predict the encoder latent)",
    "Encoder regularizer: λ ||z − sg(ẑ)||², λ = 0.1   (z stays predictable from history)",
    "Softmax curriculum: p(t) follows a sigmoid of training progress (start 0.1, end 0.7, sharpness 10)",
    "Reward-gated curriculum (info loss ∝ constant reward):",
    "  Arms once the return has risen; +0.1 when the last 3 smoothed returns stay within 10 % of the reference",
    "  −0.1 if the return falls > 20 % below it; linear floor reaches 1 at 80 % of training (deadline)",
])

tasks = sorted({r["task"] for r in runs})
rows = []
for task in tasks:
    for mode in ("curriculum", "gated", "masked", "visible"):
        rs = by[task, mode]
        if rs:
            r2 = ", ".join(f"{k.split(' R2 ')[1]} {np.mean([r['probe_r2'][k] for r in rs]):.2f}"
                           for k in rs[0]["probe_r2"] if k.startswith("ẑ"))
            rows.append([task, mode, str(len(rs)), mean_std([r["return_masked"] for r in rs]),
                         mean_std([r["return_visible"] for r in rs]), r2])
for part in (rows[: len(rows) // 2], rows[len(rows) // 2 :]):
    table("Built-in tasks: deterministic eval return (mean ± std over seeds) and ẑ probe R²",
          ["Task", "Training mode", "Seeds", "Eval masked (deploy)", "Eval visible", "ẑ R² per group"], part, size=12)
image("Eval returns per task and mode", os.path.join(D, "03_gym_returns.png"),
      "Velocities hidden (noise 0.01), mass and friction/length randomized per episode, history H = 8")
image("What the history estimator recovers: probe R² per privileged group", os.path.join(D, "04_probe_r2.png"),
      "ẑ: estimator latent (history only). z: encoder latent (privileged values). R² on held-out samples")
image("Time-based vs reward-gated masking", os.path.join(D, "05_gate_traces.png"))
if os.path.exists(os.path.join(D, "06_history_ablation.png")):
    image("Ablation: history length and observation-action history", os.path.join(D, "06_history_ablation.png"),
          "+a: previous action appended to each observation, so the history holds (o, a) pairs as in RMA")
if os.path.exists(os.path.join(D, "07_history_ablation_visible.png")):
    image("Ablation in visible mode (z keeps the dynamics): ẑ still does not recover them",
          os.path.join(D, "07_history_ablation_visible.png"),
          "Only InvertedPendulum mass rises with H = 32 (0.06 → 0.32, high seed variance); velocities gain from +a")

lam0 = [json.load(open(p)) for p in sorted(glob.glob(os.path.join(root, "ablation_lam0", "*.json")))]
if lam0:
    rows = []
    for task in sorted({r["task"] for r in lam0}):
        for mode in ("curriculum", "visible"):
            for lam, source in ((0.1, runs), (0.0, lam0)):
                rs = [r for r in source if r["task"] == task and r["mode"] == mode]
                if not rs:
                    continue
                groups = [k.split(" R2 ")[1] for k in rs[0]["probe_r2"] if k.startswith("ẑ")][1:]  # dynamics only
                r2 = lambda latent: ", ".join(  # noqa: E731
                    f"{g} {np.mean([r['probe_r2'][f'{latent} R2 {g}'] for r in rs]):.2f}" for g in groups)
                rows.append([task, mode, f"{lam:g}", mean_std([r["return_masked"] for r in rs]), r2("ẑ"), r2("z")])
    table("Encoder regularizer off (λ = 0): z keeps the dynamics, ẑ still cannot predict them",
          ["Task", "Mode", "λ", "Eval masked (deploy)", "ẑ R² dynamics", "z R² dynamics"], rows, size=11)

bullets("Takeaways", [
    "The history estimator is a good state estimator, a poor system identifier",
    "  ẑ recovers hidden velocities (R² 0.5–0.9) but not mass / friction / length (R² ≈ 0–0.1) on all tasks",
    "Why z loses the dynamics: λ||z − sg(ẑ)||² pulls z towards what the history can predict",
    "  λ = 0: z keeps mass / length (Pendulum R² 0.96, InvertedPendulum 0.9), ẑ still near 0 → the history is the bottleneck",
    "Longer history and (o, a) pairs do not fix it (H = 32 helps InvertedPendulum mass only, 0.06 → 0.32)",
    "  Low SNR: noise 0.01, short dt, dynamics affect the next state only weakly over 8–32 steps",
    "Masking curriculum: gated ≥ time-based on masked return (HalfCheetah 2675 vs 1692, Pendulum −333 vs −417)",
    "  but always-masked training is best on HalfCheetah (3429) and InvertedPendulum: little to transfer when ẑ cannot follow z",
])
bullets("Next steps", [
    "Temporal transformer estimator (attention over long (o, a) histories, as in the transformer-RL survey)",
    "Supervised system-ID head: regress e_t directly from history (auxiliary loss), not only the encoder latent",
    "Larger dynamics ranges and lower observation noise to test identifiability",
    "Tasks where dynamics matter more: Isaac Lab manipulation (object mass, friction) and locomotion (payload, terrain)",
    "Probe the policy's sensitivity: return vs true parameter, masked vs visible",
])

out =os.path.join(HERE, "History_encoder_sysid.pptx")
prs.save(out)
print(out, len(prs.slides), "slides")
