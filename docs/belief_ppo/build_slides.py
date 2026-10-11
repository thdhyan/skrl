"""Build docs/belief_ppo/BeliefPPO_pipeline.pptx from the diagrams (uploaded to Google Slides afterwards).

    .venv/bin/python docs/belief_ppo/build_slides.py
"""
import os

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
D = os.path.join(HERE, "diagrams")

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height


def title_only(title):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = title
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.4), Inches(0.2), W - Inches(0.8), Inches(0.9)
    return s


def bullets(title, items, size=20):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.4), Inches(0.2), W - Inches(0.8), Inches(1.0)
    body = s.placeholders[1]
    body.left, body.top, body.width, body.height = Inches(0.6), Inches(1.4), W - Inches(1.2), H - Inches(1.8)
    tf = body.text_frame
    tf.clear()
    for i, it in enumerate(items):
        level = 1 if it.startswith("  ") else 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.level = it.strip(), level
        p.font.size = Pt(size - 3 * level)
    return s


def image(title, png, note=None):
    s = title_only(title)
    from PIL import Image

    iw, ih = Image.open(png).size
    top = Inches(1.15)
    max_w, max_h = W - Inches(0.6), H - top - (Inches(0.6) if note else Inches(0.2))
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    s.shapes.add_picture(png, Emu((W - w) // 2), top, w, h)
    if note:
        tb = s.shapes.add_textbox(Inches(0.5), H - Inches(0.6), W - Inches(1), Inches(0.5))
        tb.text_frame.text = note
        tb.text_frame.paragraphs[0].font.size = Pt(14)
    return s


s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Belief-based PPO in skrl"
s.shapes.title.left, s.shapes.title.top = Inches(0.8), Inches(2.0)
s.shapes.title.width, s.shapes.title.height = W - Inches(1.6), Inches(1.5)
s.placeholders[1].left, s.placeholders[1].top = Inches(0.8), Inches(3.8)
s.placeholders[1].width, s.placeholders[1].height = W - Inches(1.6), Inches(1.8)
s.placeholders[1].text = (
    "How BeliefPPO changes the PPO pipeline, plus status of the Transformer-policy workstream\n"
    "branch feature/belief-pomdp-manipulation · 2026-10-10"
)

bullets("Why a belief?", [
    "POMDP: the policy cannot see part of the state (contact, object pose, velocities)",
    "Wirnshofer et al., RSS 2020: act on the belief, not on the raw observation",
    "  Particle filter, NP particles: predict(a) → update(o) → resample",
    "  Compress to an NG-component GMM: b̃ = [weights | means | log-diag-covs]",
    "  dim(b̃) = NG + 2·NG·d   (CartPole: NG=2, d=4 → 18)",
    "BeliefDQN existed already. This session adds BeliefPPO, the on-policy version",
])
image("One environment step", os.path.join(D, "01_step_pipeline.png"),
      "act() uses b̃_t from the previous step; record_transition() advances the filter to b̃_{t+1}")
image("PPO update on beliefs", os.path.join(D, "02_ppo_update.png"))
bullets("What changes compared with stock PPO", [
    "Network input: b̃ instead of the env observation; memory stores beliefs, not obs",
    "GAE bootstrap: V(b̃_{t+1}) taken from the filter after the last transition",
    "Episode reset: cfg.belief_prior_fn(o₀) re-seeds the particles on the first act() after a done",
    "Unchanged: compute_gae, clipped surrogate, value clip, KL early stop, AMP, LR scheduler (reused from PPO)",
    "Default nets: ResNet trunk (paper §V-A), Categorical policy over NC controllers + value",
    "Limit: one filter → num_envs = 1. For massively parallel Isaac Lab use pattern B",
    "Perf fix: act() reuses b̃ from record_transition() (a second GMM EM per step was ~5× slower)",
])
image("Where the filter lives", os.path.join(D, "03_patterns.png"))
image("Smoke test: CartPole POMDP, 20k steps", os.path.join(D, "04_smoke_curves.png"),
      "Mean of last 10 logs: C BeliefPPO 239 · A wrapper+PPO 466 · raw-obs baseline 53 (max 500 / 500 / 110). One seed.")
bullets("Tests and status", [
    "tests/agents/torch/test_belief_ppo.py: 8 tests (shapes, memory holds b̃, update changes params, prior reset, num_envs guard)",
    "Belief suite (filter, compression, BeliefDQN, BeliefPPO): 44 passed; transformer models: 6 passed",
    "Both belief runs reach the 500 cap; the raw-obs baseline stalls at about 100, so the belief carries the hidden velocities",
    "C is noisier than A late in training. Network and prior differ and there is one seed, so the gap is not established",
    "Next: multi-seed C vs A with the same MLP, batched BeliefPPO / pattern B for Isaac Lab",
])
image("Transformer policies instead of MLPs", os.path.join(D, "05_transformer_trunk.png"),
      "One token per observation dimension + CLS. PPO on HalfCheetah-v5, InvertedPendulum-v5, InvertedDoublePendulum-v5")
bullets("Transformer smoke runs (8 envs, 1 seed, short)", [
    "InvertedPendulum-v5, 20k steps: transformer 43 · MLP 119 (final mean return)",
    "InvertedDoublePendulum-v5, 20k steps: transformer 82",
    "HalfCheetah-v5, 50k steps: transformer −287 (still learning)",
    "Far too short to compare architectures; the transformer runs about 2× slower per step",
    "No triple-pendulum env exists in gymnasium; it needs a custom MJCF",
    "Next: ≥1M-step runs, multi-seed transformer vs MLP; feed b̃ tokens (per GMM component) to the transformer",
])

out = os.path.join(HERE, "BeliefPPO_pipeline.pptx")
prs.save(out)
print(out, len(prs.slides), "slides")
